#!/usr/bin/env python3

from __future__ import annotations

import argparse
import csv
import gzip
import os
from pathlib import Path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--samplelist", required=True, type=Path)
    parser.add_argument("--department", default="DEPT")
    parser.add_argument("--manifest-output", required=True, type=Path)
    parser.add_argument("--pedigree-output", required=True, type=Path)
    parser.add_argument("--report-output", required=True, type=Path)
    parser.add_argument("--mapping-roots", default="")
    return parser.parse_args()


def load_openpyxl():
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit(
            "openpyxl is required for Excel mapping mode. Install it in the runtime environment."
        ) from exc
    return load_workbook


def interpretation_roots(department: str, override: str) -> tuple[Path, Path]:
    if override:
        parts = [Path(p) for p in override.split("|") if p.strip()]
        if len(parts) != 2:
            raise SystemExit("--mapping-roots must contain exactly two '|' separated paths")
        return parts[0], parts[1]
    project_root = Path(f"/faststorage/project/MomaDiagnosticSamples-{department.strip()}")
    return (
        project_root / "BACKUP" / "interpretation",
        project_root / "BACKUP" / department.strip() / "extra_samples" / "interpretation",
    )


def read_sheet_rows(samplelist: Path) -> list[list[str]]:
    load_workbook = load_openpyxl()
    workbook = load_workbook(samplelist, read_only=True, data_only=True)
    sheet = workbook.active
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        values = []
        for value in row[:7]:
            if value is None:
                values.append("")
            else:
                values.append(str(value).strip())
        if any(values):
            rows.append(values)
    workbook.close()
    return rows


def sample_ids_from_rows(rows: list[list[str]]) -> list[str]:
    ids: list[str] = []
    for row in rows:
        if len(row) < 2:
            continue
        sample_id = row[1].strip()
        if sample_id and not sample_id.startswith("#"):
            ids.append(sample_id)
    return ids


def find_vcf(sample_id: str, roots: tuple[Path, Path]) -> tuple[str, Path | None]:
    for root in roots:
        base = root / sample_id
        for major in ("2", "1"):
            version_dirs = sorted([p for p in base.glob(f"v{major}.*") if p.is_dir()])
            for version_dir in version_dirs:
                vcfs = sorted(version_dir.glob("*/*.gatk.haplotype_caller.vcf.gz"))
                if vcfs:
                    return version_dir.name, vcfs[0]
    return "MISSING", None


def single_sample_name(vcf: Path) -> str:
    with gzip.open(vcf, "rt") as handle:
        for line in handle:
            if line.startswith("#CHROM"):
                cols = line.rstrip("\n").split("\t")
                samples = cols[9:]
                if len(samples) != 1:
                    raise ValueError(f"{vcf} has {len(samples)} samples; expected 1")
                return samples[0]
    raise ValueError(f"No #CHROM line found in {vcf}")


def main() -> None:
    args = parse_args()
    roots = interpretation_roots(args.department, args.mapping_roots)

    rows = read_sheet_rows(args.samplelist)
    ids = sample_ids_from_rows(rows)

    subject_to_sample: dict[str, str] = {}
    manifest_rows: list[tuple[str, str]] = []
    report_rows: list[tuple[str, str, str, str, str]] = []

    for sample_id in ids:
        version, path = find_vcf(sample_id, roots)
        if path is None:
            report_rows.append((sample_id, version, "", "", "missing"))
            continue
        try:
            vcf_sample = single_sample_name(path)
        except Exception as exc:
            report_rows.append((sample_id, version, str(path), "", f"error:{exc}"))
            continue
        subject_to_sample[sample_id] = vcf_sample
        manifest_rows.append((vcf_sample, str(path)))
        report_rows.append((sample_id, version, str(path), vcf_sample, "resolved"))

    args.manifest_output.parent.mkdir(parents=True, exist_ok=True)
    with args.manifest_output.open("w", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["sample_id", "vcf_path"])
        writer.writerows(manifest_rows)

    with args.report_output.open("w", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["subject_id", "version", "vcf_path", "vcf_sample", "status"])
        writer.writerows(report_rows)

    with args.pedigree_output.open("w", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t")
        for row in rows[1:]:
            if not row or not row[0]:
                continue
            padded = (row + [""] * 7)[:7]
            for idx in (1, 2, 3):
                value = padded[idx].strip()
                if value and value != "0":
                    padded[idx] = subject_to_sample.get(value, value)
            writer.writerow(padded)


if __name__ == "__main__":
    main()
