import sys
from pathlib import Path
from tempfile import NamedTemporaryFile

from .paths import APP_DIR, PARENT_DIR


if str(PARENT_DIR) not in sys.path:
    sys.path.insert(0, str(PARENT_DIR))

try:
    from excel_list_app import (  # noqa: E402
        SAMPLE_FILE,
        create_sample_excel,
        quote_identifier,
        read_xlsx_rows,
        safe_identifier,
        unique_headers,
    )
except ModuleNotFoundError:
    SAMPLE_FILE = APP_DIR / "data" / "sample.xlsx"

    def quote_identifier(value: str) -> str:
        return '"' + value.replace('"', '""') + '"'

    def safe_identifier(value: str, fallback: str) -> str:
        cleaned = "".join(char if char.isalnum() else "_" for char in value.strip().lower())
        cleaned = "_".join(part for part in cleaned.split("_") if part)
        return cleaned or fallback

    def unique_headers(headers: list[str]) -> list[str]:
        seen: dict[str, int] = {}
        result: list[str] = []
        for index, header in enumerate(headers):
            base = safe_identifier(str(header), f"column_{index + 1}")
            count = seen.get(base, 0)
            seen[base] = count + 1
            result.append(base if count == 0 else f"{base}_{count + 1}")
        return result

    def create_sample_excel(path: Path):
        try:
            from openpyxl import Workbook
        except ModuleNotFoundError as error:
            raise RuntimeError("openpyxl is required to create the sample workbook.") from error

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["sample", "family", "status"])
        sheet.append(["sample_001", "family_001", "affected"])
        sheet.append(["sample_002", "family_001", "parent"])
        workbook.save(path)

    def read_xlsx_rows(path: Path) -> list[list[str]]:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError as error:
            raise RuntimeError("openpyxl is required to read Excel files.") from error

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if value is None else str(value) for value in row])
        workbook.close()
        return rows


def ensure_sample_excel():
    if not SAMPLE_FILE.exists():
        create_sample_excel(SAMPLE_FILE)


def read_excel_preview(path: Path) -> tuple[list[str], list[list[str]]]:
    rows = read_xlsx_rows(path)
    if not rows:
        return [], []
    columns = [str(value) for value in rows[0]]
    data_rows = [[str(value) for value in row] for row in rows[1:]]
    return columns, data_rows


def rows_to_tsv(columns: list[str], rows: list[list[str]]) -> str:
    lines = ["\t".join(columns)]
    lines.extend("\t".join(str(value) for value in row) for row in rows)
    return "\n".join(lines)


def parse_tsv(text: str) -> tuple[list[str], list[list[str]]]:
    lines = text.splitlines()
    while lines and not lines[-1].strip():
        lines.pop()
    if not lines:
        return [], []

    columns = lines[0].split("\t")
    rows = [line.split("\t") for line in lines[1:]]
    width = max([len(columns), *(len(row) for row in rows)] or [0])
    columns = columns + [""] * (width - len(columns))
    rows = [row + [""] * (width - len(row)) for row in rows]
    return columns, rows


def write_xlsx_rows(path: Path, columns: list[str], rows: list[list[str]]):
    try:
        from openpyxl import Workbook
    except ModuleNotFoundError as error:
        raise RuntimeError("openpyxl is required to write Excel files.") from error

    path.parent.mkdir(parents=True, exist_ok=True)
    width = max([len(columns), *(len(row) for row in rows)] or [1])
    workbook = Workbook()
    sheet = workbook.active
    sheet.append((columns + [""] * (width - len(columns)))[:width])
    for row in rows:
        sheet.append((row + [""] * (width - len(row)))[:width])

    with NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temp:
        temp_path = Path(temp.name)
    try:
        workbook.save(temp_path)
        temp_path.replace(path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
